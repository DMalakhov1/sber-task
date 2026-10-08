"""Экспорт проверенного нормализованного снимка в Excel."""
import argparse
from calendar import monthrange
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

from .aggregation import SYSTEMS as ENERGY_SYSTEMS
SHEET_NAMES = ("Генерация (МВт·ч)", "Потребление (МВт·ч)", "Среднее ИБР (руб. за МВт·ч)")

def previous_month(today=None):
    today = today or datetime.now(ZoneInfo("Europe/Moscow")).date()
    last = today.replace(day=1) - timedelta(days=1)
    return last.replace(day=1), last

def parse_month(value):
    try:
        start = datetime.strptime(value, "%Y-%m").date()
        if start.strftime("%Y-%m") != value:
            raise ValueError
    except ValueError as exc:
        raise argparse.ArgumentTypeError("Нужен месяц в формате YYYY-MM") from exc
    return start, date(start.year, start.month, monthrange(start.year, start.month)[1])

def export(snapshot, output, start, end, weight, allow_partial=False, overwrite=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils import get_column_letter
    from .source import load_snapshot
    from .aggregation import validate, weighted_mean, complete_sum, quality_report, published_price_summary
    from .formulas import CachedWorkbook, strict_sum, strict_weighted, available_formulas

    rows, meta, digest = load_snapshot(snapshot)
    if meta['weight_basis'] != weight:
        raise ValueError('Выбранный вес не совпадает с методикой снимка')
    grid = validate(rows, start, end, weight, allow_partial=allow_partial)
    quality = quality_report(grid)
    wb = Workbook(); wb.remove(wb.active)
    cache = {}
    def formula(ws, cell, expression, value):
        ws[cell] = expression
        cache[ws.title, cell] = value
    fields = ('generation_mwh', 'consumption_mwh', 'ibr_rub_mwh')
    titles = ('Генерация (МВт*ч)', 'Потребление (МВт*ч)', 'Среднее ИБР (руб./МВт*ч)')
    last_data = 4 + ((end-start).days+1)*24
    month_row = last_data+1
    weight_sheet = SHEET_NAMES[0 if weight == 'generation_mwh' else 1]
    for name, title, field in zip(SHEET_NAMES, titles, fields):
        is_price = field == 'ibr_rub_mwh'
        last_col = 'N' if is_price else 'J'
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.append([title]); ws.merge_cells(f'A1:{last_col}1')
        period = f'{start:%d.%m.%Y} — {end:%d.%m.%Y}. Europe/Moscow. Факт.'
        missing_count = sum(r[field] is None for r in grid.values())
        status = ('НЕПОЛНЫЕ ДАННЫЕ: ' + str(missing_count) + ' значений отсутствуют. ') if missing_count else 'ДАННЫЕ ПОЛНЫЕ. '
        if is_price and not quality['missing_intervals'] and quality['unpublished_prices']:
            status = f'НЕПОЛНЫЕ ДАННЫЕ ИБР: {len(quality["unpublished_prices"])} цен не опубликованы источником. Интервалы загружены. '
        if meta.get('synthetic'): status = 'СИНТЕТИЧЕСКИЙ ПРИМЕР — НЕ ФАКТИЧЕСКИЕ ДАННЫЕ. ' + status
        ws.append([status + period]); ws.merge_cells(f'A2:{last_col}2')
        note = ('Вес: фактическое ' + ('потребление' if weight=='consumption_mwh' else 'генерирование' if weight=='generation_mwh' else 'объём балансирования') +
                '. ИБР = Σ(цена × вес) / Σ(вес), аналитическое допущение. J — все 7 ОЭС; K:N — только опубликованные цены, иной охват.'
                if is_price else 'Энергия за час, МВт·ч. ВСЕГО — сумма 7 ОЭС. Итоги с пропусками не рассчитываются.')
        ws.append([note]); ws.merge_cells(f'A3:{last_col}3')
        ws.append(['Дата','Час начала',*ENERGY_SYSTEMS,'Все 7 ОЭС' if is_price else 'ВСЕГО'] +
                  (['По опубликованным: взвешенное','По опубликованным: простое','Цен, шт.','Покрытие веса'] if is_price else []))
        def weights_range(c1, c2, r1, r2):
            if weight == 'balance_volume_mwh':
                # Explicit source weights remain visible to the right of the price table.
                return f'{get_column_letter(c1+13)}{r1}:{get_column_letter(c2+13)}{r2}'
            return f"'{weight_sheet}'!{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
        if is_price and weight == 'balance_volume_mwh':
            for i, system in enumerate(ENERGY_SYSTEMS,16): ws.cell(4,i,f'Вес МВт·ч: {system}')
        day = start
        while day <= end:
            for hour in range(24):
                records = [grid[day.isoformat(),hour,system] for system in ENERGY_SYSTEMS]
                values = [r[field] for r in records]
                ws.append([day,hour,*values]); n=ws.max_row
                if is_price:
                    if weight == 'balance_volume_mwh':
                        for i,record in enumerate(records,16): ws.cell(n,i,record.get(weight))
                    prices=f'C{n}:I{n}'; weights=weights_range(3,9,n,n)
                    formula(ws,f'J{n}',strict_weighted(prices,weights,7),weighted_mean((r[field],r.get(weight)) for r in records))
                    summary=published_price_summary(records,weight)
                    for col,expression,key in zip('KLMN',available_formulas(prices,weights,7),('weighted','arithmetic','published_count','weight_coverage')):
                        formula(ws,f'{col}{n}',expression,summary[key])
                else:
                    formula(ws,f'J{n}',strict_sum(f'C{n}:I{n}',7),complete_sum(values))
                for i,record in enumerate(records,3):
                    if record[field] is None:
                        key='ibr_missing_reason' if is_price else field.removesuffix('_mwh')+'_missing_reason'
                        ws.cell(n,i).comment=Comment(record.get(key,'Нет значения'),'Источник')
            day += timedelta(days=1)
        ws.cell(month_row,1,'Среднее за месяц' if is_price else 'Итого за месяц')
        for i,system in enumerate(ENERGY_SYSTEMS,3):
            records=[r for k,r in grid.items() if k[2]==system]
            col=get_column_letter(i); prices=f'{col}5:{col}{last_data}'
            value=weighted_mean((r[field],r.get(weight)) for r in records) if is_price else complete_sum(r[field] for r in records)
            expression=strict_weighted(prices,weights_range(i,i,5,last_data),len(records)) if is_price else strict_sum(prices,len(records))
            formula(ws,f'{col}{month_row}',expression,value)
        if is_price:
            prices=f'C5:I{last_data}'; weights=weights_range(3,9,5,last_data)
            formula(ws,f'J{month_row}',strict_weighted(prices,weights,len(grid)),weighted_mean((r[field],r.get(weight)) for r in grid.values()))
            summary=published_price_summary(grid.values(),weight)
            for col,expression,key in zip('KLMN',available_formulas(prices,weights,len(grid)),('weighted','arithmetic','published_count','weight_coverage')):
                formula(ws,f'{col}{month_row}',expression,summary[key])
        else:
            formula(ws,f'J{month_row}',strict_sum(f'C{month_row}:I{month_row}',7),complete_sum(r[field] for r in grid.values()))
        ws.freeze_panes='C5'
        ws.auto_filter.ref=f'A4:{last_col}{last_data}'
        ws.print_title_rows='1:4'
        ws.print_options.horizontalCentered=True
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.print_area=f'A1:{last_col}{month_row}'
        ws.column_dimensions['A'].width=24; ws.column_dimensions['B'].width=12
        for col in 'CDEFGHIJ': ws.column_dimensions[col].width=15
        for col in 'KL': ws.column_dimensions[col].width=23
        ws.column_dimensions['M'].width=11; ws.column_dimensions['N'].width=17
        for row in ws.iter_rows(min_row=1,max_row=month_row,max_col=14 if is_price else 10):
            for cell in row:
                cell.font=Font(name='Arial',size=10,color='202C38')
                cell.alignment=Alignment(vertical='center',horizontal='right' if cell.column>1 else 'left')
                if cell.row>=5 and cell.column>=3: cell.number_format='#,##0.00;[Red](#,##0.00);0.00'
            if 5<=row[0].row<=last_data: row[0].number_format='yyyy-mm-dd'
        if is_price:
            for n in range(5,month_row+1):
                ws[f'M{n}'].number_format='0';ws[f'N{n}'].number_format='0.0%'
            ws['K4'].comment=Comment('Только пары с опубликованной ценой и известным весом. Это не показатель всех 7 ОЭС.','Методика')
            ws['L4'].comment=Comment('Простое среднее опубликованных цен. За месяц — по исходным парам час × ОЭС, не среднее месячных средних.','Методика')
            ws['N4'].comment=Comment('Вес пар с опубликованной ценой / вес всех пар. При неизвестном общем весе значение пустое.','Методика')
        ws['A1'].font=Font(name='Arial',size=14,bold=True,color='165B55')
        for index in (4,month_row):
            for cell in ws[index]:
                cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor='165B55')
        for cell in ws[4]:cell.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center')
        for index,height in ((1,26),(2,38),(3,42),(4,48),(month_row,26)):
            ws.row_dimensions[index].height=height
        for n in (2,3):ws[f'A{n}'].alignment=Alignment(wrap_text=True,vertical='center')
        ws.conditional_formatting.add(f'C5:I{last_data}',FormulaRule(formula=['ISBLANK(C5)'],fill=PatternFill('solid',fgColor='FFF1D6')))
        ws['A1'].comment=Comment('Источник: '+meta['source_url']+'\nSHA-256 снимка: '+digest,'Происхождение')
    report=dict(metadata=meta,snapshot_sha256=digest,records=len(grid),start=str(start),end=str(end),
        observed_records=quality['observed_records'],complete=quality['complete'],missing_prices=len(quality['missing_prices']),
        intervals_complete=quality['intervals_complete'],volumes_complete=quality['volumes_complete'],
        prices_complete=quality['prices_complete'],published_price_summary=published_price_summary(grid.values(),weight),
        aggregation='J: all seven systems, strict; K:N: explicitly partial published-price scope; weighted sums from original pairs',
        formulas='Excel formulas with numeric caches; Excel recalculates after input edits')
    from .artifacts import save_bundle
    save_bundle(CachedWorkbook(wb,cache),output,report,quality,overwrite=overwrite)
    return report


def main():
    from pathlib import Path
    parser = argparse.ArgumentParser(description='Excel из проверенного снимка; загрузка с сайта: task3_excel.source_http')
    parser.add_argument('--month', type=parse_month, help='YYYY-MM; по умолчанию предыдущий месяц по Москве')
    parser.add_argument('--input', type=Path, required=True, help='Нормализованный JSON-снимок, см. README')
    parser.add_argument('--weight', choices=['consumption_mwh', 'generation_mwh', 'balance_volume_mwh'], required=True,
                        help='Обязательно согласовать методику; автоматического выбора нет')
    parser.add_argument('--output', type=Path)
    parser.add_argument('--allow-partial', action='store_true', help='Явно пометить неполный файл; пропуски и затронутые итоги пустые')
    parser.add_argument('--overwrite', action='store_true', help='Явно заменить готовый набор файлов с сохранением прежнего в backup')
    parser.add_argument('--check', action='store_true', help='Проверить один день без записи Excel')
    parser.add_argument('--date', type=date.fromisoformat, help='День для --check, YYYY-MM-DD')
    args = parser.parse_args()
    start, end = args.month or previous_month()
    if args.date and args.month and not start <= args.date <= end:
        parser.error('--date находится вне указанного --month')
    if args.date and not args.check:
        parser.error('--date используется только с --check')
    output = args.output or Path(__file__).parent / 'outputs' / f'electricity_{start:%Y-%m}.xlsx'
    try:
        if args.check:
            import json
            from .source import load_snapshot
            from .aggregation import validate, quality_report
            rows, meta, _ = load_snapshot(args.input)
            if meta['weight_basis'] != args.weight:
                raise ValueError('Выбранный вес не совпадает с методикой снимка')
            day = args.date or start
            selected = [r for r in rows if r['date'] == day.isoformat()]
            report = quality_report(validate(selected, day, day, args.weight, allow_partial=True))
            print(json.dumps(report, ensure_ascii=False, indent=2))
            return 0 if report['complete'] else 3
        report = export(args.input, output, start, end, args.weight, args.allow_partial, args.overwrite)
    except (ValueError, KeyError, TypeError, OSError) as exc:
        parser.exit(2, f'Ошибка: {exc}\n')
    print(f'Создан {output}; строк источника: {report["observed_records"]}/{report["records"]}; пропусков ИБР: {report["missing_prices"]}')
    return 0 if report['complete'] else 3

if __name__ == '__main__':
    raise SystemExit(main())

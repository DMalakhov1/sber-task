import json
from datetime import date
import pytest
from openpyxl import load_workbook
from task3_excel.aggregation import SYSTEMS, validate, weighted_mean
from task3_excel.export_excel import export


def snapshot():
    return {'metadata': dict(source_url='https://br.so-ups.ru/', downloaded_at='2026-10-06', timezone='Europe/Moscow', hour_convention='interval_start_0_23', energy_unit='MWh', price_unit='RUB/MWh', weight_basis='consumption_mwh', weight_explanation='СИНТЕТИЧЕСКИЙ ТЕСТ'), 'rows': [dict(date=f'2026-09-{d:02}', hour=h, system=s, status='actual', generation_mwh=10, consumption_mwh=i+1, ibr_rub_mwh=100*(i+1)) for d in range(1,31) for h in range(24) for i,s in enumerate(SYSTEMS)]}


def test_weighted():
    assert weighted_mean([(100,1),(200,3)]) == 175
    assert weighted_mean([(100,0)]) is None
    assert weighted_mean([(None,1),(200,3)]) is None
    assert weighted_mean([(0,2)]) == 0
    assert weighted_mean([(-20,2),(10,1)]) == -10
    with pytest.raises(ValueError): weighted_mean([(100,-1)])


@pytest.mark.parametrize('issue', ['missing','duplicate','forecast','nan','negative','hour','system','missing_price'])
def test_invalid(issue):
    rows = snapshot()['rows']
    if issue == 'missing': rows.pop()
    elif issue == 'duplicate': rows.append(rows[0].copy())
    elif issue == 'forecast': rows[0]['status']='forecast'
    elif issue == 'nan': rows[0]['generation_mwh']=float('nan')
    elif issue == 'negative': rows[0]['consumption_mwh']=-1
    elif issue == 'hour': rows[0]['hour']=24
    elif issue == 'system': rows[0]['system']='UNKNOWN'
    else: rows[0]['ibr_rub_mwh']=None
    with pytest.raises(ValueError): validate(rows,date(2026,9,1),date(2026,9,30),'consumption_mwh')


def test_workbook(tmp_path):
    source=tmp_path/'test.json'; source.write_text(json.dumps(snapshot()))
    output=tmp_path/'test.xlsx'
    report=export(source,output,date(2026,9,1),date(2026,9,30),'consumption_mwh')
    wb=load_workbook(output,data_only=True)
    assert len(wb.worksheets)==3 and report['records']==5040
    assert wb.worksheets[0].max_row==725
    assert wb.worksheets[0]['J725'].value==50400
    assert wb.worksheets[1]['J725'].value==20160
    assert wb.worksheets[2]['J5'].value==500
    assert wb.worksheets[2]['J725'].value==500
    with pytest.raises(ValueError): export(source,output,date(2026,9,1),date(2026,9,30),'consumption_mwh')
    original = output.read_bytes()
    export(source,output,date(2026,9,1),date(2026,9,30),'consumption_mwh',overwrite=True)
    backups = list(tmp_path.glob('test.backup-*'))
    assert len(backups) == 1 and (backups[0]/'test.xlsx').read_bytes() == original


def test_missing_price(tmp_path):
    data=snapshot(); data['rows'][0].update(ibr_rub_mwh=None,ibr_missing_reason='Не опубликовано')
    source=tmp_path/'test.json'; source.write_text(json.dumps(data)); output=tmp_path/'test.xlsx'
    with pytest.raises(ValueError, match='--allow-partial'):
        export(source,output,date(2026,9,1),date(2026,9,30),'consumption_mwh')
    export(source,output,date(2026,9,1),date(2026,9,30),'consumption_mwh',allow_partial=True)
    ws=load_workbook(output,data_only=True).worksheets[2]
    assert ws['C5'].value is None and ws['C5'].comment
    assert ws['J5'].value is None and ws['J725'].value is None
    assert ws['D725'].value==200

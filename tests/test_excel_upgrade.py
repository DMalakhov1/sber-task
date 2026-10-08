import errno
import hashlib
import json
from datetime import date
from pathlib import Path
import pytest
from openpyxl import load_workbook
from task3_excel.aggregation import published_price_summary, validate, quality_report
from task3_excel.artifacts import publish_exclusive
from task3_excel.export_excel import export
from tests.test_excel import snapshot


def test_partial_scope_uses_matching_weights_and_not_average_of_averages():
    rows=[dict(ibr_rub_mwh=p,consumption_mwh=w) for p,w in [(100,1),(200,3),(None,1000)]]
    s=published_price_summary(rows,'consumption_mwh')
    assert s['weighted']==175 and s['arithmetic']==150 and s['published_count']==2
    assert s['weight_coverage']==pytest.approx(4/1004)
    rows.append(dict(ibr_rub_mwh=400,consumption_mwh=100))
    assert published_price_summary(rows,'consumption_mwh')['weighted']==pytest.approx(40700/104)


def test_zero_weight_unknown_weight_and_negative_prices():
    def record(p,w):return dict(ibr_rub_mwh=p,consumption_mwh=w)
    s=published_price_summary([record(-10,0),record(0,0)],'consumption_mwh')
    assert s['weighted'] is None and s['weight_coverage'] is None and s['arithmetic']==-5
    s=published_price_summary([record(-10,1),record(100,None)],'consumption_mwh')
    assert s['weighted']==-10 and s['weight_coverage'] is None


def test_price_unpublished_does_not_mark_energy_incomplete(tmp_path):
    data=snapshot()
    for r in data['rows']:
        if r['system']=='Восток':r.update(ibr_rub_mwh=None,ibr_missing_reason='Источник не опубликовал ИБР')
    src=tmp_path/'x.json';src.write_text(json.dumps(data));dst=tmp_path/'x.xlsx'
    report=export(src,dst,date(2026,9,1),date(2026,9,30),'consumption_mwh',True)
    q=json.loads(dst.with_suffix('.quality.json').read_text())
    assert q['intervals_complete'] and q['volumes_complete'] and not q['prices_complete']
    assert len(q['unpublished_prices'])==720 and not q['missing_intervals']
    book=load_workbook(dst,data_only=True);ws=book.worksheets[2]
    assert 'ДАННЫЕ ПОЛНЫЕ' in book.worksheets[0]['A2'].value
    assert 'НЕПОЛНЫЕ' in ws['A2'].value
    assert ws['J5'].value is None and ws['K5'].value==pytest.approx(9100/21)
    assert ws['L5'].value==350 and ws['M5'].value==6 and ws['N5'].value==.75
    assert ws['K725'].value==pytest.approx(9100/21) and ws['M725'].value==4320
    formulas=load_workbook(dst,data_only=False)
    assert 'SUMPRODUCT' in formulas.worksheets[2]['K5'].value
    assert formulas.worksheets[0]['J5'].value.startswith('=IF(COUNT(')
    assert ws.auto_filter.ref=='A4:N724' and ws.freeze_panes=='C5'
    assert report['xlsx_sha256']==hashlib.sha256(dst.read_bytes()).hexdigest()==q['xlsx_sha256']


def test_absent_interval_distinguished_from_unpublished_price():
    rows=snapshot()['rows'];rows.pop(0)
    q=quality_report(validate(rows,date(2026,9,1),date(2026,9,30),'consumption_mwh',True))
    assert len(q['missing_intervals'])==1 and q['unpublished_prices']==[]
    assert not q['volumes_complete']


def test_unsupported_hardlinks_fall_back_without_overwrite(tmp_path,monkeypatch):
    source=tmp_path/'source';source.write_bytes(b'complete workbook')
    dest=tmp_path/'out'
    def unsupported(*a):raise OSError(errno.EOPNOTSUPP,'unsupported')
    monkeypatch.setattr('task3_excel.artifacts.os.link',unsupported)
    publish_exclusive(source,dest)
    assert dest.read_bytes()==source.read_bytes()
    with pytest.raises(FileExistsError):publish_exclusive(source,dest)
    assert dest.read_bytes()==b'complete workbook'


def test_copy_failure_removes_only_own_incomplete_file(tmp_path,monkeypatch):
    source=tmp_path/'source';source.write_bytes(b'data');dest=tmp_path/'out'
    def unsupported(*a):raise OSError(errno.EOPNOTSUPP,'unsupported')
    def fail(src,out):out.write(b'partial');raise OSError('disk failure')
    monkeypatch.setattr('task3_excel.artifacts.os.link',unsupported)
    monkeypatch.setattr('task3_excel.artifacts.shutil.copyfileobj',fail)
    with pytest.raises(OSError,match='disk failure'):publish_exclusive(source,dest)
    assert not dest.exists() and source.read_bytes()==b'data'

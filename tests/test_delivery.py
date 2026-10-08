import json
from pathlib import Path
import pytest
from openpyxl import Workbook
from task3_excel.artifacts import save_bundle
from task3_excel.source import load_snapshot


def test_export_bundle_rolls_back_io_failure(tmp_path, monkeypatch):
    import task3_excel.artifacts as module
    original = module.os.link
    count = 0
    def failing(source, target):
        nonlocal count
        count += 1
        if count == 2: raise OSError('Disk failure')
        return original(source, target)
    monkeypatch.setattr(module.os, 'link', failing)
    with pytest.raises(OSError): save_bundle(Workbook(), tmp_path/'result.xlsx', {}, {})
    assert list(tmp_path.iterdir()) == []


def test_export_bundle_does_not_replace_competing_writer(tmp_path, monkeypatch):
    import task3_excel.artifacts as module
    original = module.os.link
    def racing(source, target):
        if str(target).endswith('.xlsx'): Path(target).write_bytes(b'other writer')
        return original(source, target)
    monkeypatch.setattr(module.os, 'link', racing)
    with pytest.raises(FileExistsError): save_bundle(Workbook(), tmp_path/'result.xlsx', {}, {})
    assert (tmp_path/'result.xlsx').read_bytes() == b'other writer'
    assert len(list(tmp_path.iterdir())) == 1


@pytest.mark.parametrize('raw', ['[]','{"metadata":[],"rows":[]}', '{"metadata":{},"metadata":{}}', '{"metadata":{"a":NaN}}'])
def test_snapshot_rejects_ambiguous_or_malformed_json(tmp_path, raw):
    path=tmp_path/'source.json'; path.write_text(raw)
    with pytest.raises(ValueError): load_snapshot(path)


def test_no_output_on_wrong_extension(tmp_path):
    with pytest.raises(ValueError): save_bundle(Workbook(),tmp_path/'file.csv',{}, {})
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize('vector', [[0.,0.], [float('nan'),1.], [1.,2.,3.]])
def test_query_vectors_are_checked(tmp_path, vector):
    import numpy as np
    from task2_bot.rag.build_index import save_index
    from task2_bot.rag.search import Retriever
    chunks=[dict(id='a',source_id='s',text='A sufficiently long evidence fragment.',title='Report',url='https://example.org',page=1)]
    save_index(chunks,tmp_path,vectors=np.array([[1.,0.]]),model='fixture')
    class Fake:
        def query(self,text): return vector
    retriever=Retriever(tmp_path,embedder=Fake())
    with pytest.raises(ValueError): retriever.search('question')
